"""
Los seriales en los reportes.

Antes la columna "N.º de serial" decía "4 unidades" y ahí terminaba: el
reporte sabía cuántos aparatos había, no cuáles. Para un conteo físico eso no
alcanza — la persona tiene el equipo en la mano y necesita encontrar *ese*
número en la hoja.

Ahora:

  - **Existencias** trae, debajo de cada producto que lleva serial, una fila
    por unidad que siga en bodega, con su serial, con qué boleta entró y
    cuándo. Solo las que están: una unidad vendida no es existencia.
  - **Movimientos** trae una fila por unidad movida, que es donde se rastrea
    la que ya salió.
  - El **kardex** lleva los seriales en su fila, sin partirla: la columna de
    saldo es del movimiento entero y partirla la dejaría sin sentido.

La columna "Fila" del Excel no es adorno. La hoja no trae fila de totales
—la suma la hace quien la abre— y sin poder separar producto de unidad,
seleccionar "Valor total" contaría todo dos veces.
"""

import io

import openpyxl
from django.test import TestCase
from django.urls import reverse

from core import reportes
from core.models import Bodega
from usuarios.models import Usuario
from ventas.models import Articulo, MovimientoVenta, ingresar_unidades, sacar_unidades


class BaseReporteUnidades(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.bodega = Bodega.objects.create(nombre='Bodega 1', tipo=Bodega.Tipo.VENTA)
        cls.admin = Usuario.objects.create_user(
            username='admin_rep_uni', password='clave-de-prueba',
            rol=Usuario.Rol.ADMINISTRADOR,
        )
        cls.equipo = Articulo.objects.create(
            nombre_producto='INDICADOR SE7581P', modelo='SE7581P',
            bodega=cls.bodega, precio=1000, lleva_serie=True,
        )
        cls.repuesto = Articulo.objects.create(
            nombre_producto='CONECTOR RJ45', modelo='C-1', bodega=cls.bodega, precio=25,
        )

        cls.ingreso = MovimientoVenta.objects.create(
            folio='ING-00001', articulo=cls.equipo, cantidad=3,
            tipo_documento=MovimientoVenta.TipoDocumento.INGRESO,
            tipo_transaccion=MovimientoVenta.TipoTransaccion.VENTA,
            usuario=cls.admin,
        )
        ingresar_unidades(cls.ingreso, ['A-1001', 'A-1002', 'A-1003'])

        MovimientoVenta.objects.create(
            folio='ING-00002', articulo=cls.repuesto, cantidad=500,
            tipo_documento=MovimientoVenta.TipoDocumento.INGRESO,
            tipo_transaccion=MovimientoVenta.TipoTransaccion.REPUESTOS,
            usuario=cls.admin,
        )

    def setUp(self):
        self.client.force_login(self.admin)

    def sacar(self, serial):
        from ventas.models import UnidadArticulo

        movimiento = MovimientoVenta.objects.create(
            folio='SAL-00001', articulo=self.equipo, cantidad=1,
            tipo_documento=MovimientoVenta.TipoDocumento.SALIDA,
            tipo_transaccion=MovimientoVenta.TipoTransaccion.VENTA,
            usuario=self.admin,
        )
        sacar_unidades(movimiento, [UnidadArticulo.objects.get(numero_serie=serial)])
        return movimiento

    def hoja(self, url, **parametros):
        parametros['formato'] = 'excel'
        respuesta = self.client.get(url, parametros)
        self.assertEqual(respuesta.status_code, 200)
        return openpyxl.load_workbook(io.BytesIO(respuesta.content)).active

    def titulos(self, hoja):
        """(fila del encabezado, {título: número de columna})."""
        for fila in range(1, 15):
            for columna in range(1, 30):
                if hoja.cell(row=fila, column=columna).value == 'Fecha' or \
                        hoja.cell(row=fila, column=columna).value == 'Fila':
                    mapa = {}
                    for c in range(1, 30):
                        valor = hoja.cell(row=fila, column=c).value
                        if valor:
                            mapa[valor] = c
                    return fila, mapa
        self.fail('no se encontró el encabezado de la hoja')

    def datos(self, hoja):
        """Las filas de datos como diccionarios {título: valor}."""
        encabezado, mapa = self.titulos(hoja)
        filas = []
        for numero in range(encabezado + 1, hoja.max_row + 1):
            filas.append({
                titulo: hoja.cell(row=numero, column=columna).value
                for titulo, columna in mapa.items()
            })
        return filas


class ExistenciasDesglosaLasUnidadesTests(BaseReporteUnidades):
    def test_debajo_del_producto_van_sus_seriales(self):
        filas = self.datos(self.hoja(reverse('reporte_existencias')))

        del_equipo = [f for f in filas if f['Código'] == self.equipo.codigo_interno]
        self.assertEqual(
            [(f['Fila'], f['N.º de serial']) for f in del_equipo],
            [('Producto', '3 unidades'),
             ('Unidad', 'A-1001'), ('Unidad', 'A-1002'), ('Unidad', 'A-1003')],
        )

    def test_un_repuesto_sigue_siendo_una_sola_fila(self):
        filas = self.datos(self.hoja(reverse('reporte_existencias')))

        del_repuesto = [f for f in filas if f['Código'] == self.repuesto.codigo_interno]
        self.assertEqual(len(del_repuesto), 1)
        self.assertEqual(del_repuesto[0]['N.º de serial'], 'S/S')

    def test_la_unidad_dice_con_que_boleta_entro_y_cuando(self):
        filas = self.datos(self.hoja(reverse('reporte_existencias')))

        unidad = next(f for f in filas if f['N.º de serial'] == 'A-1001')
        self.assertEqual(unidad['Entró con'], 'ING-00001')
        self.assertIsNotNone(unidad['Fecha de ingreso'])

    def test_la_carga_inicial_se_dice_con_su_nombre(self):
        """El ajuste de saldo no lleva boleta porque no hubo papel."""
        otro = Articulo.objects.create(
            nombre_producto='OTRO EQUIPO', modelo='X-9', bodega=self.bodega,
            precio=50, lleva_serie=True,
        )
        inicial = MovimientoVenta.objects.create(
            articulo=otro, cantidad=1,
            tipo_documento=MovimientoVenta.TipoDocumento.INGRESO,
            tipo_transaccion=MovimientoVenta.TipoTransaccion.AJUSTE_INICIAL,
            usuario=self.admin,
        )
        ingresar_unidades(inicial, ['B-1'])

        filas = self.datos(self.hoja(reverse('reporte_existencias')))

        unidad = next(f for f in filas if f['N.º de serial'] == 'B-1')
        self.assertEqual(unidad['Entró con'], 'Carga inicial')

    def test_la_que_ya_salio_no_aparece(self):
        """
        Es un reporte de existencias: dice lo que hay. La unidad vendida se
        rastrea en el de movimientos.
        """
        self.sacar('A-1002')

        filas = self.datos(self.hoja(reverse('reporte_existencias')))

        seriales = [f['N.º de serial'] for f in filas]
        self.assertNotIn('A-1002', seriales)
        self.assertIn('A-1001', seriales)

    def test_el_nivel_va_solo_en_la_fila_del_producto(self):
        """
        Una unidad sola tiene existencia 1 contra un umbral de 2: si la
        columna se calculara por fila, las tres dirían "Crítico" de un
        producto que tiene tres.
        """
        filas = self.datos(self.hoja(reverse('reporte_existencias')))
        del_equipo = [f for f in filas if f['Código'] == self.equipo.codigo_interno]

        producto, unidades = del_equipo[0], del_equipo[1:]
        self.assertTrue(producto['Nivel'])
        self.assertEqual([u['Nivel'] for u in unidades], [None, None, None])

    def test_las_cantidades_no_se_cuentan_dos_veces_al_filtrar(self):
        """
        La razón de la columna "Fila": la hoja no trae totales, y sumar sin
        separar contaría el producto y otra vez sus unidades.
        """
        filas = self.datos(self.hoja(reverse('reporte_existencias')))

        productos = [f for f in filas if f['Fila'] == 'Producto']
        unidades = [f for f in filas if f['Fila'] == 'Unidad']

        self.assertEqual(sum(f['Existencia'] for f in productos), 503)
        self.assertEqual(sum(f['Existencia'] for f in unidades), 3)
        self.assertEqual(sum(f['Valor total'] for f in productos), 3 * 1000 + 500 * 25)

    def test_cada_unidad_vale_el_precio_del_producto(self):
        filas = self.datos(self.hoja(reverse('reporte_existencias')))

        unidad = next(f for f in filas if f['N.º de serial'] == 'A-1001')
        self.assertEqual(unidad['Existencia'], 1)
        self.assertEqual(unidad['Precio unitario'], 1000)
        self.assertEqual(unidad['Valor total'], 1000)


class LaPantallaDeExistenciasTests(BaseReporteUnidades):
    def test_muestra_el_resumen_y_deja_desplegar(self):
        respuesta = self.client.get(reverse('reporte_existencias'))

        self.assertContains(respuesta, '3 unidades')
        self.assertContains(respuesta, 'desplegar-unidades')

    def test_las_filas_de_unidad_vienen_escondidas(self):
        respuesta = self.client.get(reverse('reporte_existencias'))

        self.assertContains(respuesta, 'A-1001')
        self.assertContains(respuesta, 'fila-unidad')
        self.assertRegex(respuesta.content.decode(), r'class="fila-unidad"[^>]*hidden')

    def test_un_repuesto_no_ofrece_desplegar(self):
        solo_repuesto = self.client.get(
            reverse('reporte_existencias'), {'bodega': self.bodega.pk},
        )
        html = solo_repuesto.content.decode()
        fila = html[html.index(self.repuesto.codigo_interno):][:400]

        self.assertNotIn('desplegar-unidades', fila)


class MovimientosDesglosaLasUnidadesTests(BaseReporteUnidades):
    def test_una_fila_por_unidad_movida(self):
        filas = self.datos(self.hoja(reverse('reporte_movimientos')))

        del_ingreso = [f for f in filas if f['Boleta'] == 'ING-00001']
        self.assertEqual(
            sorted(f['N.º de serial'] for f in del_ingreso),
            ['A-1001', 'A-1002', 'A-1003'],
        )
        self.assertEqual([f['Cantidad'] for f in del_ingreso], [1, 1, 1])

    def test_ahi_si_aparece_la_que_ya_salio(self):
        """Lo que se dijo del reporte de existencias tiene que ser cierto."""
        self.sacar('A-1002')

        filas = self.datos(self.hoja(reverse('reporte_movimientos')))

        salidas = [f for f in filas if f['Boleta'] == 'SAL-00001']
        self.assertEqual([f['N.º de serial'] for f in salidas], ['A-1002'])

    def test_un_repuesto_sigue_en_una_fila_con_su_cantidad(self):
        filas = self.datos(self.hoja(reverse('reporte_movimientos')))

        del_repuesto = [f for f in filas if f['Boleta'] == 'ING-00002']
        self.assertEqual(len(del_repuesto), 1)
        self.assertEqual(del_repuesto[0]['N.º de serial'], 'S/S')
        self.assertEqual(del_repuesto[0]['Cantidad'], 500)

    def test_la_pantalla_dice_lo_mismo(self):
        respuesta = self.client.get(reverse('reporte_movimientos'))

        self.assertContains(respuesta, 'A-1001')
        self.assertContains(respuesta, 'N.º de serial')


class ElKardexTests(BaseReporteUnidades):
    def test_lleva_los_seriales_del_movimiento(self):
        respuesta = self.client.get(reverse('kardex_articulo', args=[self.equipo.pk]))

        for serial in ('A-1001', 'A-1002', 'A-1003'):
            self.assertContains(respuesta, serial)

    def test_no_parte_la_fila_para_no_romper_el_saldo(self):
        """
        El saldo es del movimiento entero. Con una fila por unidad, la columna
        diría tres veces el mismo saldo o tres saldos inventados.
        """
        respuesta = self.client.get(reverse('kardex_articulo', args=[self.equipo.pk]))

        self.assertEqual(len(respuesta.context['movimientos'].object_list), 1)

    def test_la_fila_entera_lleva_a_la_boleta(self):
        """
        Como en el catálogo, en Entradas y salidas y en la ficha. Era la única
        tabla del sistema donde había que apuntarle al número.
        """
        respuesta = self.client.get(reverse('kardex_articulo', args=[self.equipo.pk]))

        self.assertContains(
            respuesta,
            f'data-href="{reverse("documento_detalle", args=["ING-00001"])}"',
        )

    def test_un_repuesto_dice_ss(self):
        respuesta = self.client.get(reverse('kardex_articulo', args=[self.repuesto.pk]))

        self.assertContains(respuesta, 'S/S')


class ElDesgloseEnPythonTests(BaseReporteUnidades):
    """Sobre las funciones, sin pasar por una pantalla."""

    def test_existencias_da_producto_mas_sus_unidades(self):
        _filas, detalle, _totales = reportes.existencias()
        filas = reportes.desglosar_existencias(detalle)

        del_equipo = [f for f in filas if f.articulo == self.equipo]
        self.assertEqual([f.tipo for f in del_equipo],
                         ['Producto', 'Unidad', 'Unidad', 'Unidad'])

    def test_la_fila_de_producto_conserva_la_existencia_real(self):
        _filas, detalle, _totales = reportes.existencias()
        filas = reportes.desglosar_existencias(detalle)

        producto = next(f for f in filas if f.articulo == self.equipo and f.es_producto)
        self.assertEqual(producto.existencia, 3)
        self.assertEqual(producto.valor, 3000)

    def test_movimientos_reparte_por_serial(self):
        _resumen, detalle = reportes.movimientos_del_periodo()
        filas = reportes.desglosar_movimientos(detalle)

        del_ingreso = [f for f in filas if f.folio == 'ING-00001']
        self.assertEqual(len(del_ingreso), 3)
        self.assertEqual([f.cantidad for f in del_ingreso], [1, 1, 1])

    def test_la_fila_de_movimiento_deja_llegar_a_lo_del_movimiento(self):
        """Delega en el movimiento: fecha, folio, producto, cliente."""
        _resumen, detalle = reportes.movimientos_del_periodo()
        fila = reportes.desglosar_movimientos(detalle)[0]

        self.assertEqual(fila.articulo, fila.movimiento.articulo)
        self.assertEqual(fila.folio, fila.movimiento.folio)

    def test_preguntar_por_algo_que_no_existe_no_se_cuelga(self):
        """
        Regresión posible: __getattr__ que busca en `movimiento` sin guardia
        se llama a sí mismo para siempre si el atributo tampoco está ahí.
        """
        _resumen, detalle = reportes.movimientos_del_periodo()
        fila = reportes.desglosar_movimientos(detalle)[0]

        with self.assertRaises(AttributeError):
            fila.esto_no_existe
